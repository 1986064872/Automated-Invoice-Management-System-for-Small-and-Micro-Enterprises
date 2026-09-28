"""端到端自检：对着真实运行的后端，走一遍项目书第十四章的验收清单。

用法（后端已在 127.0.0.1:8000 运行时）：
    python backend/scripts/e2e_check.py

样本从哪来
----------
用的是**脱敏公开版**样本 `backend/testdata/测试发票公开版/`，不是真实票据。
两个好处：
  ① 仓库里不含任何真实企业名/税号/票号 —— 脚本里的期望值全是假数据，可以放心提交；
  ② 别人 clone 下来就能直接跑（真实票据在 `素材/` 下且已被 gitignore，clone 后压根不存在）。
"""

from __future__ import annotations

import json
import os
import sys
import time
from pathlib import Path

import httpx

# 默认打本机 8000；可以用 E2E_BASE 指向别的实例
# （配合后端的 APP_DATA_DIR 就能在一个临时库上跑，完全不碰真实数据）
BASE = os.environ.get("E2E_BASE", "http://127.0.0.1:8000").rstrip("/") + "/api/v1"
PROJECT_DIR = Path(__file__).resolve().parents[2]
# 脱敏公开版样本（已进仓库）
SAMPLE_DIR = Path(__file__).resolve().parents[1] / "testdata" / "测试发票公开版"

PASS = "✔"
FAIL = "✘"
results: list[tuple[bool, str]] = []


def find_sample(name: str) -> Path | None:
    """素材目录可以自由整理（含子文件夹），这里按文件名递归找。"""
    if not SAMPLE_DIR.exists():
        return None
    for path in SAMPLE_DIR.rglob(name):
        if path.is_file():
            return path
    return None


def check(ok: bool, label: str, extra: str = "") -> None:
    results.append((ok, label))
    mark = PASS if ok else FAIL
    print(f"  [{mark}] {label}" + (f" —— {extra}" if extra else ""))


def wait_job(client: httpx.Client, job_id: str, timeout: float = 90.0) -> dict:
    deadline = time.time() + timeout
    job = {}
    while time.time() < deadline:
        resp = client.get(f"{BASE}/jobs/{job_id}", params={"with_files": True})
        resp.raise_for_status()
        job = resp.json()
        if job["status"] in ("done", "failed", "partial"):
            return job
        time.sleep(0.4)
    return job


# 本脚本自己上传的测试票据，按原文件名前缀识别
# （公开版文件名本身就是脱敏产物；与真实票据的对应关系只存在 gitignored 的
#   anonymize_map.local.json 里，不写进仓库）
TEST_FILE_PREFIXES = (
    "09_26322000006800000009",   # 票A（PDF）
    "14_26322000006800000014",   # 票B（PDF）
    "02_26322000006200000002",   # 票C（PDF）
    "图片_08",                    # 图片票
    "重复票副本",
)


def purge(client: httpx.Client, purge_all: bool) -> int:
    """清理库里已有的票据。

    ⚠️ 安全护栏（血泪教训）：
        删除接口会**连磁盘上的原票文件一起删**，而且是不可逆的（早期版本直接 unlink）。
        所以默认**只删本脚本自己上传的测试票据**；一旦发现有别人的数据，
        立刻中止，绝不默默清库。
        确实要清空全部，显式加 --purge-all。

    返回：删除条数；返回 -1 表示「发现真实数据，已中止」。
    """
    listing = client.get(f"{BASE}/invoices", params={"page_size": 200}).json()
    items = listing.get("items", [])
    if not items:
        return 0

    def is_test_item(item: dict) -> bool:
        name = (item.get("file") or {}).get("original_name") or ""
        return name.startswith(TEST_FILE_PREFIXES)

    foreign = [i for i in items if not is_test_item(i)]
    if foreign and not purge_all:
        print(f"\n  [中止] 库里发现 {len(foreign)} 条**非测试**数据，不会自动清理：")
        for item in foreign[:8]:
            name = (item.get("file") or {}).get("original_name") or "?"
            print(f"        · {item.get('invoice_number') or '(无号码)'}  {name}")
        if len(foreign) > 8:
            print(f"        … 还有 {len(foreign) - 8} 条")
        print("  这是保护你的真实票据（删除会连原票文件一起删）。")
        print()
        print("  注意：这里刻意「发现一条非测试数据就整体中止」，不做部分清理 ——")
        print("       部分清理既救不了真实数据，又会让本次验收跑在半脏的库上。")
        print()
        print("  想在这个库上跑验收，请二选一：")
        print("    ① 换一个空数据库跑（推荐）：设 APP_DB_PATH 指向临时文件后重启后端")
        print("    ② 确实要清空全部：python backend/scripts/e2e_check.py --purge-all")
        print()
        return -1

    # 走到这里说明库里全是本脚本自己传的测试数据，可以放心清
    targets = items
    removed = 0
    for item in targets:
        client.delete(f"{BASE}/invoices/{item['id']}", params={"delete_file": True})
        removed += 1
    return removed


def main() -> int:
    purge_all = "--purge-all" in sys.argv
    client = httpx.Client(timeout=120.0)

    print("=" * 78)
    print("0) 健康检查")
    health = client.get(f"{BASE}/health").json()
    check(health.get("ok") is True, "后端存活", f"OCR 供应商 = {health.get('ocr_provider')}")

    removed = purge(client, purge_all)
    if removed < 0:
        print("=" * 78)
        print("已中止：请先备份或转移真实票据，或用 --purge-all 明确表示要清空。")
        print("=" * 78)
        return 1
    if removed:
        print(f"  （已清理上一轮遗留的 {removed} 条测试票据，保证本次验收是干净环境）")

    print("\n1) 上传混合批次（3 张 PDF + 1 张重复副本 + 1 张图片 + 1 个非法文件）")
    targets = [
        find_sample("09_26322000006800000009.pdf"),
        find_sample("14_26322000006800000014.pdf"),
        find_sample("02_26322000006200000002.pdf"),
        find_sample("图片_08.jpg"),
    ]
    missing = [name for name, path in zip(
        ["票A", "票B", "票C", "图片票"], targets
    ) if path is None]
    if missing:
        print("  在公开版样本目录里没找到：", missing)
        print(f"  （目录：{SAMPLE_DIR}）")
        print("  先跑一次：python backend/scripts/make_public_samples.py")
        return 1
    targets = [t for t in targets if t is not None]

    files = []
    handles = []
    for path in targets:
        fh = path.open("rb")
        handles.append(fh)
        files.append(("files", (path.name, fh, "application/octet-stream")))
    # 重复票：同一张 PDF 换个名字再传一次
    dup = targets[1]
    dup_fh = dup.open("rb")
    handles.append(dup_fh)
    files.append(("files", ("重复票副本.pdf", dup_fh, "application/octet-stream")))
    # 非法格式
    files.append(("files", ("说明.txt", b"not an invoice", "text/plain")))

    resp = client.post(f"{BASE}/invoices/upload", files=files)
    for fh in handles:
        fh.close()
    resp.raise_for_status()
    upload = resp.json()

    check(len(upload["accepted"]) == 5, f"接收 5 个合法文件", f"实际 {len(upload['accepted'])}")
    check(len(upload["rejected"]) == 1, "拒绝 1 个非法格式文件",
          upload["rejected"][0]["reason"] if upload["rejected"] else "")
    job_id = upload["job_id"]

    print("\n2) 批次进度")
    job = wait_job(client, job_id)
    check(job["status"] in ("done", "partial"), "批次跑完", f"状态 = {job['status']}")
    check(job["total"] == 5, "批次文件数 = 5", f"实际 {job['total']}")
    check(job["succeeded"] == 5, "5 个文件都产出了待复核记录",
          f"成功 {job['succeeded']} / 失败 {job['failed']}")
    check(len(job["files"]) == 5 and all(f["status"] != "queued" for f in job["files"]),
          "每个文件都有独立状态（无卡在排队中）")

    print("\n3) 字段识别准确度（对照票面人工核对值）")
    listing = client.get(f"{BASE}/invoices", params={"page_size": 100}).json()
    by_number = {i["invoice_number"]: i for i in listing["items"] if i["invoice_number"]}

    expectations = {
        # 三张票的期望值全部来自脱敏公开版样本（金额沿用真实票面，因为脱敏不改金额）
        "26322000006800000009": {
            "date": "2026-08-20", "seller": "aa贸易有限公司",
            "total": 19250.00, "tax": 2214.60, "untaxed": 17035.40,
        },
        "26322000006800000014": {
            "date": "2026-08-20", "seller": "aa贸易有限公司",
            "total": 12600.00, "tax": 1449.56, "untaxed": 11150.44,
        },
        "26322000006200000002": {
            "date": "2026-07-29", "seller": "jj贸易有限公司",
            "total": 25500.00, "tax": 2933.63, "untaxed": 22566.37,
        },
    }
    for number, expected in expectations.items():
        inv = by_number.get(number)
        if inv is None:
            check(False, f"识别到发票 {number}")
            continue
        ok = (
            inv["invoice_date"] == expected["date"]
            and inv["seller_name"] == expected["seller"]
            and abs((inv["total_amount"] or 0) - expected["total"]) < 0.01
            and abs((inv["tax_amount"] or 0) - expected["tax"]) < 0.01
            and abs((inv["amount_without_tax"] or 0) - expected["untaxed"]) < 0.01
        )
        check(
            ok,
            f"发票 {number} 号码/日期/销方/金额/税额/价税合计全对",
            f"{inv['invoice_date']} | {inv['seller_name']} | "
            f"{inv['amount_without_tax']}+{inv['tax_amount']}={inv['total_amount']}",
        )

    print("\n4) 图片票：本地 PaddleOCR（GPU）自动识别，不丢原文件")
    image_invoices = [
        i for i in listing["items"] if i["file"] and i["file"]["file_type"] == "jpg"
    ]
    check(len(image_invoices) == 1, "图片票生成了待复核记录")
    if image_invoices:
        img = image_invoices[0]
        check(
            img["file"]["ocr_provider"] == "paddleocr",
            "识别引擎是本地 PaddleOCR",
            str(img["file"]["ocr_provider"]),
        )
        # 图片票的期望值同样来自公开版样本（图片_08.jpg）
        expect = {
            "invoice_number": "23942000000000000001",
            "invoice_date": "2023-11-01",
            "seller_name": "mm贸易有限公司",
            "buyer_name": "nn贸易有限公司",
            "buyer_tax_id": "987654321987654331",
            "total_amount": 100.0,
            "amount_without_tax": 99.01,
            "tax_amount": 0.99,
        }
        bad = []
        for field, want in expect.items():
            got = img.get(field)
            if isinstance(want, float):
                ok = got is not None and abs(got - want) < 0.01
            else:
                ok = got == want
            if not ok:
                bad.append(f"{field}: 期望 {want} 实得 {got}")
        check(not bad, "图片票关键字段全部识别正确", "; ".join(bad))
        check(len(img["items"]) >= 1, "图片票明细行已解析",
              f"{len(img['items'])} 行")
        check(
            not any(f["code"] == "manual_entry_required" for f in img["risk_flags"]),
            "不再需要人工录入",
        )
        check(img["error_count"] == 0, "无红色错误，可直接入账")

    print("\n5) 重复识别（相同发票再传一次 → 红色错误拦截）")
    dup_invoices = [
        i for i in listing["items"]
        if i["file"] and i["file"]["original_name"] == "重复票副本.pdf"
    ]
    check(len(dup_invoices) == 1, "重复副本已入库待复核")
    if dup_invoices:
        dup_inv = dup_invoices[0]
        codes = [f["code"] for f in dup_inv["risk_flags"]]
        check("duplicate_invoice" in codes, "标记为疑似重复（红色错误）", ",".join(codes))
        check(dup_inv["error_count"] >= 1, "error_count >= 1")

        resp = client.post(f"{BASE}/invoices/{dup_inv['id']}/confirm", json={"confirmed_by": "e2e"})
        check(resp.status_code == 409, "确认入账被拒绝（409）", f"HTTP {resp.status_code}")

    print("\n6) 人工复核 → 确认入账")
    confirmable = [
        i for i in listing["items"]
        if i["error_count"] == 0 and i["status"] == "pending_review" and i["invoice_number"]
    ]
    check(len(confirmable) >= 3, f"有 {len(confirmable)} 张干净票据可确认")
    confirmed_ids = []
    for inv in confirmable:
        resp = client.post(f"{BASE}/invoices/{inv['id']}/confirm", json={"confirmed_by": "e2e"})
        if resp.status_code == 200:
            confirmed_ids.append(inv["id"])
        else:
            print("     确认失败：", resp.status_code, resp.text[:200])
    check(len(confirmed_ids) == len(confirmable), f"全部 {len(confirmable)} 张确认入账成功")

    # 幂等：再确认一次不应产生第二条账目
    if confirmed_ids:
        first = confirmed_ids[0]
        client.post(f"{BASE}/invoices/{first}/confirm", json={"confirmed_by": "e2e"})
        again = client.get(f"{BASE}/invoices/{first}").json()
        check(again["status"] == "confirmed", "重复确认是幂等的（状态仍为已入账）")

    print("\n7) 账本筛选")
    ledger = client.get(f"{BASE}/ledger", params={"status": "confirmed"}).json()
    check(ledger["total"] == len(confirmed_ids), "已入账条数与确认数一致",
          f"账本 {ledger['total']} / 确认 {len(confirmed_ids)}")
    check(ledger["summary"]["total_amount"] > 0, "汇总金额正常",
          f"价税合计合计 ¥{ledger['summary']['total_amount']}")

    month = "2026-08"
    month_ledger = client.get(f"{BASE}/ledger", params={"month": month}).json()
    check(month_ledger["total"] > 0, f"按月份 {month} 筛选可用",
          f"{month_ledger['total']} 条 / ¥{month_ledger['summary']['total_amount']}")

    print("\n8) Excel 导出（按明细行展开，含单位/数量/单价）")
    payload = {"month": month, "status": "all"}
    preview = client.get(f"{BASE}/exports/preview", params=payload).json()
    resp = client.post(f"{BASE}/exports/excel", json=payload)
    check(resp.status_code == 200, "导出接口返回文件", f"HTTP {resp.status_code}")
    if resp.status_code == 200:
        out = PROJECT_DIR / "data" / "e2e_export.xlsx"
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(resp.content)
        check(out.stat().st_size > 4000, "xlsx 文件已生成", f"{out.stat().st_size} 字节")

        from openpyxl import load_workbook

        wb = load_workbook(out)
        ws = wb["费用明细"]
        header = [c.value for c in ws[1]]
        expected_header = [
            "开票日期", "发票类型", "发票代码", "发票号码", "销售方名称", "销售方税号",
            "项目名称", "规格型号", "单位", "数量", "单价",
            "不含税金额", "税额", "价税合计", "费用分类", "会计科目", "状态", "原文件名", "备注",
        ]
        check(header == expected_header, "表头 19 列完全符合约定（含单位/数量/单价）",
              " / ".join(str(h) for h in header[6:11]))
        check(ws.freeze_panes == "A2", "首行已冻结")
        check(ws.auto_filter.ref is not None, "已启用筛选", str(ws.auto_filter.ref))
        # 数据驱动地核对：导出应该等于「本月发票展开成明细行」的结果
        matched = [i for i in listing["items"] if (i["invoice_date"] or "").startswith(month)]
        expected_rows = sum(max(1, len(i["items"])) for i in matched)
        expected_amount = round(sum((i["amount_without_tax"] or 0) for i in matched), 2)

        rows = list(ws.iter_rows(min_row=2, values_only=True))
        check(len(rows) == expected_rows, "导出行数 == 本月发票展开的明细行数",
              f"表格 {len(rows)} / 期望 {expected_rows}（{len(matched)} 张票）")
        check(preview["invoice_count"] == len(matched), "预览的发票数正确",
              f"{preview['invoice_count']} vs {len(matched)}")

        units = {r[8] for r in rows}
        check(units == {"台"}, "单位列取到了票面单位", str(units))

        quantities = sorted(r[9] for r in rows if r[9] is not None)
        expected_qty = sorted(
            it["quantity"]
            for i in matched
            for it in (i["items"] or [{"quantity": None}])
            if it.get("quantity") is not None
        )
        check(quantities == expected_qty, "数量列与票面一致", str(quantities))

        check(all(r[10] is not None for r in rows), "单价列有值",
              "、".join(str(round(r[10], 4)) for r in rows))

        item_sum = round(sum((r[11] or 0) for r in rows), 2)
        check(abs(item_sum - expected_amount) < 0.01,
              "明细行金额相加 == 对应发票票面不含税合计",
              f"{item_sum} vs {expected_amount}")

    print("\n8b) 账本多选：勾选导出 + 批量删除")
    # 拿两张本次上传的票来试（隔离环境里全是测试数据；跑生产库时 purge 会先拦住）
    picked = [i for i in listing["items"] if i.get("items")]
    if len(picked) >= 2:
        pick_ids = [picked[0]["id"], picked[1]["id"]]
        # 用「原文件名」而不是发票号码来比对 —— 测试票的号码可能是空的，
        # 拿号码比对会变成 {None} == {None}，看着通过其实什么都没验到。
        pick_files = {
            (p.get("file") or {}).get("original_name") for p in picked[:2]
        }

        sel = client.post(f"{BASE}/exports/excel", json={"status": "all", "invoice_ids": pick_ids})
        check(sel.status_code == 200, "勾选导出返回文件", f"HTTP {sel.status_code}")
        if sel.status_code == 200:
            from openpyxl import load_workbook
            import io as _io

            ws_sel = load_workbook(_io.BytesIO(sel.content))["费用明细"]
            head = [c.value for c in ws_sel[1]]
            name_col = head.index("原文件名") + 1
            got_files = {ws_sel.cell(row=r, column=name_col).value for r in range(2, ws_sel.max_row + 1)}
            check(got_files == pick_files,
                  "勾选导出只含选中的那几张票（不夹带其他票）",
                  f"导出 {sorted(got_files)} / 勾选 {sorted(pick_files)}")
            expect_rows = sum(max(1, len(p["items"])) for p in picked[:2])
            check(ws_sel.max_row - 1 == expect_rows, "勾选导出的行数 == 选中票的明细行数",
                  f"{ws_sel.max_row - 1} vs {expect_rows}")

        # 安全边界：传空列表必须是「一张都不导」，绝不能理解成「不过滤 → 导全部」
        empty = client.post(f"{BASE}/exports/excel", json={"status": "all", "invoice_ids": []})
        check(empty.status_code != 200, "空勾选列表被拒绝（不会误导出全部）",
              f"HTTP {empty.status_code}")

        before_total = client.get(f"{BASE}/invoices", params={"page_size": 1}).json()["total"]
        batch = client.post(f"{BASE}/invoices/batch-delete",
                            json={"ids": pick_ids, "delete_file": True})
        check(batch.status_code == 200, "批量删除接口可用", f"HTTP {batch.status_code}")
        body = batch.json() if batch.status_code == 200 else {}
        check(body.get("deleted") == len(pick_ids), "批量删除的条数正确",
              f"deleted={body.get('deleted')}")
        check(not body.get("failed"), "批量删除没有失败项", str(body.get("failed")))
        after_total = client.get(f"{BASE}/invoices", params={"page_size": 1}).json()["total"]
        check(after_total == before_total - len(pick_ids), "删除后总数相应减少",
              f"{before_total} → {after_total}")
        still = client.get(f"{BASE}/invoices", params={"page_size": 200}).json()["items"]
        check(all(i["id"] not in pick_ids for i in still), "选中的票确实删掉了")

        # 不存在的 id 不能把整批拖崩，要如实报告
        ghost = client.post(f"{BASE}/invoices/batch-delete",
                            json={"ids": ["not-a-real-id"], "delete_file": False})
        gbody = ghost.json() if ghost.status_code == 200 else {}
        check(ghost.status_code == 200 and gbody.get("failed"),
              "批量删除遇到不存在的 id 会单独报失败（不崩整批）",
              f"HTTP {ghost.status_code} / {str(gbody.get('failed'))[:60]}")
    else:
        check(False, "账本多选：需要至少两张带明细的票据才能验", f"当前 {len(picked)} 张")

    print("\n9) 分类规则引擎")
    rules = client.get(f"{BASE}/rules").json()
    check(len(rules) > 0, "系统内置规则已初始化", f"{len(rules)} 条")
    test = client.get(f"{BASE}/rules/test", params={"item_name": "*电子工业设备*LED液晶电视机"}).json()
    check(test["expense_category"] != "待分类", "规则命中测试可用",
          f"{test['expense_category']} ← {test['rule_source']}")

    print("\n10) 概览")
    dash = client.get(f"{BASE}/dashboard/summary").json()
    check("month" in dash and "recent_jobs" in dash, "Dashboard 接口正常",
          f"{dash['month']}：票据 {dash['invoice_count']} 张 / 金额 ¥{dash['total_amount']} / "
          f"待复核 {dash['pending_review']} / 异常 {dash['abnormal']}")

    # ---------------- 汇总 ----------------
    total = len(results)
    passed = sum(1 for ok, _ in results if ok)
    print("\n" + "=" * 78)
    print(f"验收结果：{passed}/{total} 通过")
    if passed < total:
        print("未通过项：")
        for ok, label in results:
            if not ok:
                print("  -", label)
    print("=" * 78)
    return 0 if passed == total else 2


if __name__ == "__main__":
    sys.exit(main())
